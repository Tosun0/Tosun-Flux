from __future__ import annotations

import csv
import json
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from PIL import Image, ImageOps

try:
    from pypdf import PdfReader, PdfWriter
except ImportError:  # pragma: no cover - packaging/runtime guard
    PdfReader = PdfWriter = None

try:
    from docx import Document
except ImportError:  # pragma: no cover - packaging/runtime guard
    Document = None


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".ico"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".wmv", ".flv", ".m4v", ".gif"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".wma"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json"}


class ConversionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ConversionResult:
    source: Path
    outputs: tuple[Path, ...]


@dataclass(frozen=True)
class ConversionOptions:
    optimize: str = "source"
    resolution: str = "source"
    aspect: str = "source"
    fit: str = "fit"
    width: int | None = None
    height: int | None = None
    fps: str = "source"


RESOLUTIONS = {
    "4k-uhd": (3840, 2160),
    "4k": (4096, 2160),
    "qhd": (2560, 1440),
    "fhd": (1920, 1080),
    "hd": (1280, 720),
    "sd": (720, 480),
}

ASPECTS = {
    "16:9": 16 / 9,
    "9:16": 9 / 16,
    "1:1": 1.0,
    "4:3": 4 / 3,
    "3:4": 3 / 4,
}


def app_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
    return Path(__file__).resolve().parent


def bundled_tool(name: str, system_name: str | None = None) -> Path | None:
    candidates = [app_root() / "vendor" / name]
    found = shutil.which(system_name or name)
    if found:
        candidates.append(Path(found))
    return next((candidate for candidate in candidates if candidate.exists()), None)


def source_metadata(path: Path) -> dict[str, object]:
    metadata: dict[str, object] = {"path": str(path), "kind": file_kind(path), "width": None, "height": None, "fps": None, "duration": None}
    try:
        if metadata["kind"] == "image":
            with Image.open(path) as image:
                metadata["width"], metadata["height"] = image.size
            return metadata

        if metadata["kind"] != "video":
            return metadata
        tool = bundled_tool("ffmpeg.exe", "ffmpeg")
        if tool is None:
            return metadata
        completed = subprocess.run(
            [str(tool), "-hide_banner", "-i", str(path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        probe = completed.stderr
        size_match = re.search(r"Video:.*?(\d{2,6})x(\d{2,6})", probe, re.IGNORECASE | re.DOTALL)
        fps_match = re.search(r"(\d+(?:\.\d+)?)\s+(?:fps|tbr)", probe, re.IGNORECASE)
        duration_match = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", probe, re.IGNORECASE)
        if size_match:
            metadata["width"], metadata["height"] = int(size_match.group(1)), int(size_match.group(2))
        if fps_match:
            metadata["fps"] = fps_match.group(1)
        if duration_match:
            hours, minutes, seconds = duration_match.groups()
            metadata["duration"] = int(hours) * 3600 + int(minutes) * 60 + float(seconds)
    except (OSError, ValueError):
        pass
    return metadata


def file_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix in AUDIO_EXTENSIONS:
        return "audio"
    if suffix in VIDEO_EXTENSIONS and suffix != ".gif":
        return "video"
    if suffix in IMAGE_EXTENSIONS:
        return "image"
    if suffix in TEXT_EXTENSIONS:
        return "text"
    return "unknown"


def supported_targets(path: Path) -> tuple[str, ...]:
    kind = file_kind(path)
    if kind == "image":
        return ("png", "jpg", "webp", "bmp", "tiff", "gif", "pdf")
    if kind == "pdf":
        return ("png", "jpg", "pdf")
    if kind == "docx":
        return ("txt", "md")
    if kind == "text":
        suffix = path.suffix.lower()
        if suffix in {".csv", ".tsv"}:
            return ("csv", "json", "txt")
        if suffix == ".json":
            return ("json", "csv", "txt")
        return ("txt", "md", "docx")
    if kind == "video":
        return ("mp4", "webm", "mov", "mkv", "avi", "gif", "png-sequence", "jpg-sequence")
    if kind == "audio":
        return ("mp3", "wav", "flac", "m4a", "ogg")
    return ()


def common_targets(paths: Iterable[Path]) -> tuple[str, ...]:
    items = list(paths)
    if not items:
        return ()
    common = set(supported_targets(items[0]))
    for path in items[1:]:
        common &= set(supported_targets(path))
    order = ("png", "jpg", "webp", "bmp", "tiff", "gif", "pdf", "txt", "md", "csv", "json", "docx", "mp4", "webm", "mov", "mkv", "avi", "png-sequence", "jpg-sequence", "mp3", "wav", "flac", "m4a", "ogg")
    return tuple(target for target in order if target in common)


def unique_output(directory: Path, stem: str, extension: str) -> Path:
    candidate = directory / f"{stem}.{extension}"
    index = 1
    while candidate.exists():
        candidate = directory / f"{stem} ({index}).{extension}"
        index += 1
    return candidate


def unique_sequence_pattern(directory: Path, stem: str, extension: str) -> tuple[Path, str]:
    index = 0
    while True:
        sequence_stem = stem if index == 0 else f"{stem} ({index})"
        if not any(directory.glob(f"{sequence_stem}_*.{extension}")):
            return directory / f"{sequence_stem}_%06d.{extension}", sequence_stem
        index += 1


def _is_identity_conversion(source: Path, target: str, options: ConversionOptions) -> bool:
    source_format = {"jpeg": "jpg", "tif": "tiff", "markdown": "md"}.get(source.suffix.lower().lstrip("."), source.suffix.lower().lstrip("."))
    return (
        source_format == target.lower()
        and options.optimize == "source"
        and options.resolution == "source"
        and options.aspect == "source"
        and options.width is None
        and options.height is None
        and options.fps == "source"
    )


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ConversionError(f"텍스트 인코딩을 읽을 수 없습니다: {path.name}")


def target_dimensions(source_size: tuple[int, int], options: ConversionOptions) -> tuple[int, int] | None:
    if options.width is not None and options.height is not None:
        return options.width, options.height
    if options.resolution == "source" and options.aspect == "source":
        return None

    source_width, source_height = source_size
    if options.resolution in RESOLUTIONS:
        base_width, base_height = RESOLUTIONS[options.resolution]
        long_edge = base_width
    else:
        base_width, base_height = source_width, source_height
        long_edge = max(source_width, source_height)

    if options.aspect in ASPECTS:
        ratio = ASPECTS[options.aspect]
        if ratio >= 1:
            width, height = long_edge, round(long_edge / ratio)
        else:
            height, width = long_edge, round(long_edge * ratio)
        return max(2, width // 2 * 2), max(2, height // 2 * 2)

    ratio = source_width / source_height
    scale = min(base_width / source_width, base_height / source_height)
    width = round(source_width * scale)
    height = round(source_height * scale)
    return max(2, width // 2 * 2), max(2, height // 2 * 2)


def _apply_image_geometry(image: Image.Image, target: str, options: ConversionOptions) -> Image.Image:
    size = target_dimensions(image.size, options)
    if size is None or size == image.size:
        return image.copy()
    if options.fit == "stretch":
        return image.resize(size, Image.Resampling.LANCZOS)
    if options.fit == "fill":
        return ImageOps.fit(image, size, Image.Resampling.LANCZOS)

    contained = ImageOps.contain(image, size, Image.Resampling.LANCZOS)
    transparent = target in {"png", "webp", "tiff"} and "A" in image.getbands()
    mode = "RGBA" if transparent else "RGB"
    canvas = Image.new(mode, size, (0, 0, 0, 0) if transparent else "black")
    if contained.mode != mode:
        contained = contained.convert(mode)
    canvas.paste(contained, ((size[0] - contained.width) // 2, (size[1] - contained.height) // 2), contained if transparent else None)
    contained.close()
    return canvas


def _write_image(image: Image.Image, destination: Path, target: str, options: ConversionOptions) -> None:
    target_format = {"jpg": "JPEG", "tiff": "TIFF"}.get(target, target.upper())
    converted = image
    if target == "jpg" and image.mode not in {"RGB", "L"}:
        background = Image.new("RGB", image.size, "white")
        if "A" in image.getbands():
            background.paste(image, mask=image.getchannel("A"))
        else:
            background.paste(image)
        converted = background
    if target == "pdf" and image.mode not in {"RGB", "L"}:
        converted = image.convert("RGB")
    save_options: dict[str, object] = {}
    quality = {"quality": 92, "balanced": 82, "small": 68}.get(options.optimize)
    if options.optimize == "source" and target == "jpg":
        save_options.update(quality=95, subsampling=0)
    elif options.optimize == "source" and target == "webp":
        save_options.update(lossless=True)
    elif quality and target in {"jpg", "webp"}:
        save_options.update(quality=quality, optimize=True)
    elif options.optimize != "source" and target == "png":
        save_options.update(optimize=True, compress_level={"quality": 6, "balanced": 8, "small": 9}[options.optimize])
    converted.save(destination, format=target_format, **save_options)
    if converted is not image:
        converted.close()


def _convert_image(source: Path, output_dir: Path, target: str, options: ConversionOptions) -> ConversionResult:
    destination = unique_output(output_dir, source.stem, target)
    with Image.open(source) as image:
        transformed = _apply_image_geometry(image, target, options)
        try:
            _write_image(transformed, destination, target, options)
        finally:
            transformed.close()
    return ConversionResult(source, (destination,))


def _convert_pdf(source: Path, output_dir: Path, target: str, options: ConversionOptions) -> ConversionResult:
    if target == "pdf":
        if PdfReader is None or PdfWriter is None:
            raise ConversionError("PDF 압축 모듈을 찾을 수 없습니다.")
        destination = unique_output(output_dir, source.stem, target)
        if options.optimize == "source":
            shutil.copy2(source, destination)
            return ConversionResult(source, (destination,))
        reader = PdfReader(source)
        writer = PdfWriter(clone_from=reader)
        quality = {"quality": 90, "balanced": 78, "small": 62}[options.optimize]
        for page in writer.pages:
            page.compress_content_streams()
            for image_file in list(page.images):
                try:
                    replacement = image_file.image.convert("RGB")
                    try:
                        image_file.replace(replacement, quality=quality, optimize=True)
                    finally:
                        replacement.close()
                except TypeError:
                    pass  # Inline images cannot be replaced, but stream compression still applies.
        writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
        temporary = destination.with_suffix(".tmp.pdf")
        with temporary.open("wb") as handle:
            writer.write(handle)
        if temporary.stat().st_size < source.stat().st_size:
            temporary.replace(destination)
        else:
            temporary.unlink()
            shutil.copy2(source, destination)
        return ConversionResult(source, (destination,))

    tool = bundled_tool("pdftoppm.exe", "pdftoppm")
    if tool is None:
        raise ConversionError("PDF 렌더러(pdftoppm)를 찾을 수 없습니다.")
    output_stem = unique_output(output_dir, source.stem, target).with_suffix("")
    dpi = {"4k": 300, "qhd": 220, "fhd": 150, "hd": 110}.get(
        options.resolution,
        {"quality": 220, "balanced": 150, "small": 110}.get(options.optimize, 150),
    )
    render_format = "jpeg" if target == "jpg" else target
    args = [str(tool), "-r", str(dpi), f"-{render_format}", str(source), str(output_stem)]
    completed = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if completed.returncode:
        message = completed.stderr.strip() or "PDF 변환에 실패했습니다."
        raise ConversionError(message)
    outputs = tuple(sorted(output_dir.glob(f"{output_stem.name}-*.{target}")))
    if not outputs:
        raise ConversionError("PDF 페이지 이미지를 만들지 못했습니다.")
    if options.aspect != "source" or (options.width is not None and options.height is not None):
        for output in outputs:
            with Image.open(output) as page:
                transformed = _apply_image_geometry(page, target, options)
                try:
                    _write_image(transformed, output, target, options)
                finally:
                    transformed.close()
    return ConversionResult(source, outputs)


def _video_filter(options: ConversionOptions, source_size: tuple[int, int] = (1920, 1080)) -> str | None:
    size = target_dimensions(source_size, options)
    if size is None:
        return None
    width, height = (max(2, value // 2 * 2) for value in size)
    if options.fit == "stretch":
        return f"scale={width}:{height}:flags=lanczos"
    if options.fit == "fill":
        return f"scale={width}:{height}:force_original_aspect_ratio=increase:flags=lanczos,crop={width}:{height}"
    return f"scale={width}:{height}:force_original_aspect_ratio=decrease:flags=lanczos,pad={width}:{height}:(ow-iw)/2:(oh-ih)/2:color=black"


def _convert_media(source: Path, output_dir: Path, target: str, options: ConversionOptions) -> ConversionResult:
    tool = bundled_tool("ffmpeg.exe", "ffmpeg")
    if tool is None:
        raise ConversionError("FFmpeg를 찾을 수 없습니다.")
    args = [str(tool), "-hide_banner", "-loglevel", "error", "-y", "-i", str(source)]
    if file_kind(source) == "video":
        metadata = source_metadata(source)
        source_size = (
            int(metadata["width"]),
            int(metadata["height"]),
        ) if metadata.get("width") and metadata.get("height") else (1920, 1080)
        video_filter = _video_filter(options, source_size)
        if video_filter:
            args.extend(["-vf", video_filter])
        if options.fps != "source":
            args.extend(["-r", options.fps])
        if target in {"png-sequence", "jpg-sequence"}:
            extension = "png" if target == "png-sequence" else "jpg"
            destination, sequence_stem = unique_sequence_pattern(output_dir, source.stem, extension)
            args.extend(["-an"])
            if extension == "jpg":
                args.extend(["-q:v", {"source": "2", "quality": "2", "balanced": "4", "small": "7"}[options.optimize]])
            elif options.optimize != "source":
                args.extend(["-compression_level", {"quality": "4", "balanced": "7", "small": "9"}[options.optimize]])
            args.append(str(destination))
            completed = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace")
            if completed.returncode:
                message = completed.stderr.strip() or "프레임 시퀀스 변환에 실패했습니다."
                raise ConversionError(message)
            outputs = tuple(sorted(output_dir.glob(f"{sequence_stem}_*.{extension}")))
            if not outputs:
                raise ConversionError("프레임을 만들지 못했습니다.")
            return ConversionResult(source, outputs)
    destination = unique_output(output_dir, source.stem, target)
    if file_kind(source) == "video":
        if options.optimize != "source" and target != "gif":
            crf = {"quality": "18", "balanced": "23", "small": "28"}[options.optimize]
            if target == "webm":
                args.extend(["-c:v", "libvpx-vp9", "-crf", crf, "-b:v", "0", "-c:a", "libopus"])
            else:
                audio_rate = {"quality": "192k", "balanced": "160k", "small": "128k"}[options.optimize]
                args.extend(["-c:v", "libx264", "-preset", "medium", "-crf", crf, "-c:a", "aac", "-b:a", audio_rate])
    args.append(str(destination))
    completed = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if completed.returncode:
        message = completed.stderr.strip() or "미디어 변환에 실패했습니다."
        raise ConversionError(message)
    return ConversionResult(source, (destination,))


def _convert_docx(source: Path, output_dir: Path, target: str) -> ConversionResult:
    if Document is None:
        raise ConversionError("DOCX 변환 모듈을 찾을 수 없습니다.")
    document = Document(source)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    destination = unique_output(output_dir, source.stem, target)
    destination.write_text(text, encoding="utf-8")
    return ConversionResult(source, (destination,))


def _text_rows(source: Path) -> list[dict[str, str]]:
    delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def _convert_text(source: Path, output_dir: Path, target: str) -> ConversionResult:
    destination = unique_output(output_dir, source.stem, target)
    suffix = source.suffix.lower()
    if target == "docx":
        if Document is None:
            raise ConversionError("DOCX 변환 모듈을 찾을 수 없습니다.")
        document = Document()
        for line in _read_text(source).splitlines() or [""]:
            document.add_paragraph(line)
        document.save(destination)
        return ConversionResult(source, (destination,))
    if target in {"txt", "md"}:
        if suffix in {".csv", ".tsv"}:
            rows = _text_rows(source)
            text = "\n".join("\t".join(row.values()) for row in rows)
        elif suffix == ".json":
            text = json.dumps(json.loads(_read_text(source)), ensure_ascii=False, indent=2)
        else:
            text = _read_text(source)
        destination.write_text(text, encoding="utf-8")
        return ConversionResult(source, (destination,))
    if target == "json":
        if suffix == ".json":
            payload = json.loads(_read_text(source))
        else:
            payload = _text_rows(source)
        destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return ConversionResult(source, (destination,))
    if target == "csv":
        if suffix == ".json":
            payload = json.loads(_read_text(source))
            if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
                raise ConversionError("JSON → CSV는 객체 배열만 지원합니다.")
            rows = payload
        else:
            rows = _text_rows(source)
        fields: list[str] = []
        for row in rows:
            for field in row:
                if field not in fields:
                    fields.append(field)
        with destination.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        return ConversionResult(source, (destination,))
    raise ConversionError(f"지원하지 않는 텍스트 변환입니다: {source.suffix} → .{target}")


def convert_file(source: Path, output_dir: Path, target: str, options: ConversionOptions | None = None) -> ConversionResult:
    source = Path(source)
    output_dir = Path(output_dir)
    options = options or ConversionOptions()
    if not source.is_file():
        raise ConversionError(f"입력 파일을 찾을 수 없습니다: {source}")
    if target not in supported_targets(source):
        raise ConversionError(f"지원하지 않는 변환입니다: {source.suffix} → .{target}")
    output_dir.mkdir(parents=True, exist_ok=True)
    if _is_identity_conversion(source, target, options):
        destination = unique_output(output_dir, source.stem, target)
        shutil.copy2(source, destination)
        return ConversionResult(source, (destination,))
    kind = file_kind(source)
    if kind == "image":
        return _convert_image(source, output_dir, target, options)
    if kind == "pdf":
        return _convert_pdf(source, output_dir, target, options)
    if kind == "docx":
        return _convert_docx(source, output_dir, target)
    if kind == "text":
        return _convert_text(source, output_dir, target)
    if kind in {"video", "audio"}:
        return _convert_media(source, output_dir, target, options)
    raise ConversionError(f"지원하지 않는 파일 형식입니다: {source.suffix or '(확장자 없음)'}")


def convert_files(paths: Iterable[Path], output_dir: Path, target: str, on_progress: Callable[[int, int, ConversionResult | None, Exception | None], None] | None = None, options: ConversionOptions | None = None) -> list[ConversionResult]:
    items = list(paths)
    results: list[ConversionResult] = []
    for index, source in enumerate(items, start=1):
        try:
            result = convert_file(source, output_dir, target, options)
            results.append(result)
            if on_progress:
                on_progress(index, len(items), result, None)
        except Exception as error:
            if on_progress:
                on_progress(index, len(items), None, error)
    return results


def demo() -> None:
    assert common_targets([Path("photo.png"), Path("photo.jpg")])[0] == "png"
    assert "mp4" in supported_targets(Path("clip.mov"))
    assert "json" in supported_targets(Path("data.csv"))


if __name__ == "__main__":
    demo()
    print("converter self-check: ok")
